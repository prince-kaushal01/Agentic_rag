from pathlib import Path
import openpyxl
from ingestion.parsers.base import BaseParser, ParsedDocument


class ExcelParser(BaseParser):
    supported_extensions = [".xlsx", ".xls"]

    def parse(self, path: Path) -> ParsedDocument:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        sections = []
        all_parts = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            headers = [str(c) if c is not None else "" for c in rows[0]]
            text_parts = [f"Sheet: {sheet_name} | Columns: {', '.join(headers)}"]

            for i, row in enumerate(rows[1:], 1):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    line = " | ".join(
                        f"{h}: {v}" for h, v in zip(headers, cells) if v
                    )
                    text_parts.append(f"Row {i}: {line}")

            section_text = "\n".join(text_parts)
            sections.append({"title": sheet_name, "text": section_text})
            all_parts.append(section_text)

        wb.close()
        return ParsedDocument(
            file_name=path.name,
            file_type="xlsx",
            title=path.stem.replace("_", " ").replace("-", " "),
            content="\n\n".join(all_parts),
            sections=sections,
            metadata={"sheets": wb.sheetnames, "source_path": str(path)},
        )
